from datetime import datetime, timezone
import csv
import os

from airflow import DAG
from airflow.operators.python import PythonOperator
from airflow.providers.postgres.hooks.postgres import PostgresHook
from airflow.providers.common.sql.operators.sql import SQLExecuteQueryOperator

# ---------- Общие настройки ----------
POSTGRES_CONN_ID = "etl_postgres"   # Connection из Admin → Connections


def to_snake(name):
    # "Order Date" -> "order_date", "Sub-Category" -> "sub_category"
    return name.strip().lower().replace(" ", "_").replace("-", "_")


def utc_now_text():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


# ---------- Источник 1: CSV с заказами ----------
SOURCE_FILE = "/opt/airflow/input/superstore.csv"        # исходный файл (cp1252)
STAGE_FILE = "/opt/airflow/output/superstore_raw.csv"    # промежуточный файл (UTF-8)
SOURCE_ENCODING = "cp1252"

EXPECTED_COLUMNS = [
    "Row ID", "Order ID", "Order Date", "Ship Date", "Ship Mode",
    "Customer ID", "Customer Name", "Segment", "Country", "City",
    "State", "Postal Code", "Region", "Product ID", "Category",
    "Sub-Category", "Product Name", "Sales", "Quantity", "Discount", "Profit",
]
RAW_COLUMNS = [to_snake(c) for c in EXPECTED_COLUMNS]

# ---------- Источник 2: Excel со справочниками ----------
EXCEL_FILE = "/opt/airflow/input/superstore_extra.xlsx"
# Лист Orders не грузим: заказы уже приходят из CSV
EXCEL_SHEETS = {
    "People":  {"table": "people",  "columns": ["Person", "Region"]},
    "Returns": {"table": "returns", "columns": ["Returned", "Order ID"]},
}


def excel_stage_file(table):
    return f"/opt/airflow/output/superstore_{table}_raw.csv"


def cell_to_text(value):
    # Excel хранит числа как float: 1.0 -> "1", 42420.0 -> "42420"
    if value is None:
        return ""                      # пустое поле -> NULL при COPY
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def is_empty_row(values):
    return all(v is None or str(v).strip() == "" for v in values)


# ---------- Загрузка в PostgreSQL (общая для обеих веток) ----------
def load_table(cur, table, columns, stage_file, source_file):
    # Создать таблицу, удалить строки этого файла, загрузить заново, посчитать
    columns_ddl = ", ".join(f"{c} TEXT" for c in columns)
    cur.execute(f"""
        CREATE SCHEMA IF NOT EXISTS raw;
        CREATE TABLE IF NOT EXISTS raw.{table} (
            {columns_ddl},
            source_file TEXT,
            loaded_at TIMESTAMP
        );
    """)

    cur.execute(f"DELETE FROM raw.{table} WHERE source_file = %s", (source_file,))
    print(f"raw.{table}: deleted old rows for {source_file}: {cur.rowcount}")

    all_columns = ", ".join(columns + ["source_file", "loaded_at"])
    with open(stage_file, encoding="utf-8") as f:
        cur.copy_expert(
            f"COPY raw.{table} ({all_columns}) FROM STDIN WITH (FORMAT csv, HEADER true)", f
        )

    cur.execute(f"SELECT COUNT(*) FROM raw.{table} WHERE source_file = %s", (source_file,))
    return cur.fetchone()[0]


def run_loads_in_transaction(loads):
    # loads: список (table, columns, stage_file, source_file, expected_rows)
    # Все таблицы фиксируются вместе: всё или ничего
    hook = PostgresHook(postgres_conn_id=POSTGRES_CONN_ID)
    conn = hook.get_conn()
    try:
        with conn.cursor() as cur:
            for table, columns, stage_file, source_file, expected in loads:
                loaded = load_table(cur, table, columns, stage_file, source_file)
                if loaded != expected:
                    raise ValueError(f"raw.{table}: loaded {loaded} rows, expected {expected}")
                print(f"Loaded {loaded} rows into raw.{table}")
        conn.commit()
    except Exception:
        conn.rollback()   # при ошибке старые данные остаются нетронутыми
        raise
    finally:
        conn.close()


# ---------- Ветка 1: CSV ----------
def read_csv():
    # Проверить файл и заголовок, посчитать строки
    if not os.path.exists(SOURCE_FILE):
        raise FileNotFoundError(f"File not found: {SOURCE_FILE}")

    with open(SOURCE_FILE, encoding=SOURCE_ENCODING, newline="") as f:
        reader = csv.reader(f)
        header = next(reader)
        row_count = sum(1 for _ in reader)

    if header != EXPECTED_COLUMNS:
        raise ValueError(f"Unexpected header: {header}")

    print(f"Header OK, {len(header)} columns, {row_count} rows")
    return row_count


def transform(ti):
    # cp1252 -> UTF-8, snake_case заголовки, технические колонки
    expected_rows = ti.xcom_pull(task_ids="read_csv")
    source_name = os.path.basename(SOURCE_FILE)
    loaded_at = utc_now_text()

    rows = 0
    with open(SOURCE_FILE, encoding=SOURCE_ENCODING, newline="") as src, \
         open(STAGE_FILE, "w", encoding="utf-8", newline="") as dst:
        reader = csv.reader(src)
        writer = csv.writer(dst)

        next(reader)
        writer.writerow(RAW_COLUMNS + ["source_file", "loaded_at"])

        for line_no, row in enumerate(reader, start=2):
            if len(row) != len(EXPECTED_COLUMNS):
                raise ValueError(f"Line {line_no}: expected {len(EXPECTED_COLUMNS)} fields, got {len(row)}")
            writer.writerow(row + [source_name, loaded_at])
            rows += 1

    if rows != expected_rows:
        raise ValueError(f"Row count mismatch: read_csv={expected_rows}, transform={rows}")

    print(f"Written {rows} rows to {STAGE_FILE}, loaded_at={loaded_at} UTC")
    return {"source_file": source_name, "rows": rows}


def load_postgres(ti):
    meta = ti.xcom_pull(task_ids="transform")
    run_loads_in_transaction([
        ("superstore", RAW_COLUMNS, STAGE_FILE, meta["source_file"], meta["rows"]),
    ])


# ---------- Ветка 2: Excel ----------
def read_excel():
    # Проверить файл, листы и заголовки, посчитать непустые строки
    import openpyxl   # импорт внутри задачи: DAG разбирается быстрее

    if not os.path.exists(EXCEL_FILE):
        raise FileNotFoundError(f"File not found: {EXCEL_FILE}")

    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
    try:
        counts = {}
        for sheet, cfg in EXCEL_SHEETS.items():
            if sheet not in wb.sheetnames:
                raise ValueError(f"Sheet not found: {sheet}. Available: {wb.sheetnames}")

            n = len(cfg["columns"])
            rows = wb[sheet].iter_rows(values_only=True)
            header = [cell_to_text(v) for v in next(rows)][:n]
            if header != cfg["columns"]:
                raise ValueError(f"Sheet {sheet}: unexpected header {header}")

            total = empty = 0
            for row in rows:
                total += 1
                if is_empty_row(row[:n]):
                    empty += 1
            counts[sheet] = total - empty
            print(f"Sheet {sheet}: header OK, {counts[sheet]} rows, {empty} empty rows skipped")
    finally:
        wb.close()

    return counts


def transform_excel(ti):
    # Каждый лист -> свой UTF-8 CSV: snake_case, числа без ".0", без пустых строк
    import openpyxl

    expected = ti.xcom_pull(task_ids="read_excel")
    source_name = os.path.basename(EXCEL_FILE)
    loaded_at = utc_now_text()

    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
    try:
        tables = {}
        for sheet, cfg in EXCEL_SHEETS.items():
            n = len(cfg["columns"])
            stage_file = excel_stage_file(cfg["table"])
            rows_iter = wb[sheet].iter_rows(values_only=True)
            next(rows_iter)   # пропускаем заголовок

            written = 0
            with open(stage_file, "w", encoding="utf-8", newline="") as dst:
                writer = csv.writer(dst)
                writer.writerow([to_snake(c) for c in cfg["columns"]] + ["source_file", "loaded_at"])
                for row in rows_iter:
                    values = list(row[:n]) + [None] * (n - len(row[:n]))
                    if is_empty_row(values):
                        continue
                    writer.writerow([cell_to_text(v) for v in values] + [source_name, loaded_at])
                    written += 1

            if written != expected[sheet]:
                raise ValueError(f"Sheet {sheet}: read_excel={expected[sheet]}, transform={written}")

            tables[cfg["table"]] = {"file": stage_file, "rows": written}
            print(f"Sheet {sheet}: written {written} rows to {stage_file}")
    finally:
        wb.close()

    return {"source_file": source_name, "tables": tables}


def load_excel_postgres(ti):
    meta = ti.xcom_pull(task_ids="transform_excel")
    loads = []
    for cfg in EXCEL_SHEETS.values():
        info = meta["tables"][cfg["table"]]
        columns = [to_snake(c) for c in cfg["columns"]]
        loads.append((cfg["table"], columns, info["file"], meta["source_file"], info["rows"]))
    run_loads_in_transaction(loads)


# ---------- Проверки качества данных ----------
DQ_SQL_FILE = os.path.join(os.path.dirname(__file__), "sql", "dq_checks.sql")


def check_quality(run_id):
    # Выполнить проверки, сохранить результаты в dq.check_results,
    # упасть, если хоть одна проверка уровня ERROR не прошла
    with open(DQ_SQL_FILE, encoding="utf-8") as f:
        checks_sql = f.read()

    checked_at = utc_now_text()
    hook = PostgresHook(postgres_conn_id=POSTGRES_CONN_ID)
    conn = hook.get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(checks_sql)          # последним идёт SELECT с результатами
            results = cur.fetchall()

            # при повторном запуске (Clear) старые результаты этого run_id заменяются
            cur.execute("DELETE FROM dq.check_results WHERE run_id = %s", (run_id,))

            rows = []
            for check_no, check_name, severity, failed_count in results:
                if failed_count == 0:
                    status = "PASS"
                elif severity == "ERROR":
                    status = "FAIL"
                else:
                    status = "WARN"
                rows.append((run_id, checked_at, check_no, check_name, severity, failed_count, status))

            cur.executemany(
                "INSERT INTO dq.check_results "
                "(run_id, checked_at, check_no, check_name, severity, failed_count, status) "
                "VALUES (%s, %s, %s, %s, %s, %s, %s)",
                rows,
            )
        conn.commit()   # результаты сохраняются ДО того, как задача упадёт
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    for _, _, check_no, check_name, severity, failed_count, status in rows:
        print(f"{status:4} | #{check_no:>2} {check_name:<30} | {severity:<7} | failed: {failed_count}")

    failed = [r for r in rows if r[6] == "FAIL"]
    warned = [r for r in rows if r[6] == "WARN"]
    print(f"Summary: {len(rows)} checks, {len(failed)} FAIL, {len(warned)} WARN")

    if failed:
        names = ", ".join(f"#{r[2]} {r[3]} ({r[5]})" for r in failed)
        raise ValueError(f"Data quality check failed: {names}")

    return {"checks": len(rows), "failed": 0, "warnings": len(warned)}


# ---------- SQL-задачи ----------
def sql_task(task_id, sql_file):
    # Одна SQL-задача = один файл из dags/sql
    return SQLExecuteQueryOperator(
        task_id=task_id,
        conn_id=POSTGRES_CONN_ID,
        sql=sql_file,
        show_return_value_in_logs=True,
    )


# ---------- DAG ----------
with DAG(
    dag_id="superstore_raw_dag",
    start_date=datetime(2026, 9, 1),
    schedule=None,      # пока только ручной запуск
    catchup=False,
    max_active_runs=1,  # новый запуск ждёт окончания предыдущего
    tags=["etl", "superstore"],
) as dag:

    # Ветка 1: CSV -> raw.superstore
    t1 = PythonOperator(task_id="read_csv", python_callable=read_csv)
    t2 = PythonOperator(task_id="transform", python_callable=transform)
    t3 = PythonOperator(task_id="load_postgres", python_callable=load_postgres)

    # Ветка 2: Excel -> raw.people, raw.returns
    e1 = PythonOperator(task_id="read_excel", python_callable=read_excel)
    e2 = PythonOperator(task_id="transform_excel", python_callable=transform_excel)
    e3 = PythonOperator(task_id="load_excel_postgres", python_callable=load_excel_postgres)

    # Проверки качества: стоят между загрузкой и staging
    dq = PythonOperator(task_id="check_quality", python_callable=check_quality)

    # ---------- staging: по одной задаче на таблицу ----------
    stg_superstore = sql_task("stg_superstore", "sql/staging/superstore.sql")
    stg_people     = sql_task("stg_people",     "sql/staging/people.sql")
    stg_returns    = sql_task("stg_returns",    "sql/staging/returns.sql")

    # ---------- mart: звёздная схема ----------
    dim_customer = sql_task("dim_customer", "sql/mart/dim_customer.sql")
    dim_product  = sql_task("dim_product",  "sql/mart/dim_product.sql")
    dim_date     = sql_task("dim_date",     "sql/mart/dim_date.sql")
    dim_location = sql_task("dim_location", "sql/mart/dim_location.sql")
    fact_sales   = sql_task("fact_sales",   "sql/mart/fact_sales.sql")

    t1 >> t2 >> t3
    e1 >> e2 >> e3
    [t3, e3] >> dq                                   # проверки ждут ОБЕ загрузки
    dq >> [stg_superstore, stg_people, stg_returns]  # staging параллельно

    stg_superstore >> [dim_customer, dim_product, dim_date, dim_location]
    stg_people >> dim_location                       # адреса + менеджеры

    [dim_customer, dim_product, dim_date, dim_location, stg_returns] >> fact_sales
